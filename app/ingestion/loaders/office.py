import os

import logfire
from docx import Document
from pptx import Presentation


def _docx_text(file_path: str) -> str:
    doc = Document(file_path)
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append("\t".join(cells))
    return "\n".join(parts)


def _shape_text(shape: object) -> list[str]:
    parts: list[str] = []
    if getattr(shape, "has_text_frame", False):
        text = getattr(shape, "text_frame").text.strip()
        if text:
            parts.append(text)
    if getattr(shape, "has_table", False):
        for row in getattr(shape, "table").rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append("\t".join(cells))
    return parts


def _pptx_text(file_path: str) -> str:
    prs = Presentation(file_path)
    parts: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            parts.extend(_shape_text(shape))
    return "\n".join(parts)


def parse_office(file_path: str) -> str:
    """Extract text from a .docx or .pptx file."""
    with logfire.span("Office Document Parsing", filename=file_path):
        try:
            ext = os.path.splitext(file_path)[1].lower()
            if ext == ".docx":
                full_text = _docx_text(file_path)
            elif ext == ".pptx":
                full_text = _pptx_text(file_path)
            else:
                raise ValueError(f"Unsupported office file: {file_path}")

            if not full_text.strip():
                logfire.warning(f"No text extracted from {file_path}")
            else:
                logfire.info(f"Extracted {len(full_text)} characters from {file_path}")

            return full_text
        except Exception as e:
            logfire.error(f"Office parse failed for {file_path}: {e}")
            raise
